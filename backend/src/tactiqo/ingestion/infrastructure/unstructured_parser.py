"""Isolated Unstructured parser with approved native fallbacks."""

import csv
import tempfile
from importlib.metadata import version
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any
from uuid import UUID, uuid4

import anyio
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader
from unstructured.partition.auto import partition

from tactiqo.knowledge.domain.models import CanonicalDocumentElement


class UnstructuredDocumentParser:
    """Normalize parser output so vendor types never cross the adapter boundary."""

    name = "unstructured"
    version = version("unstructured")

    async def parse(
        self,
        document_id: UUID,
        name: str,
        content_type: str,
        content: bytes,
    ) -> list[CanonicalDocumentElement]:
        """Partition content in a worker thread and retain element provenance."""
        return await anyio.to_thread.run_sync(
            self._parse_sync,
            document_id,
            name,
            content_type,
            content,
        )

    def _parse_sync(
        self,
        document_id: UUID,
        name: str,
        content_type: str,
        content: bytes,
    ) -> list[CanonicalDocumentElement]:
        suffix = Path(name).suffix.lower() or self._suffix_for(content_type)
        with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as handle:
            handle.write(content)
            temporary_path = Path(handle.name)
        try:
            try:
                elements = partition(filename=str(temporary_path), strategy="fast")
                normalized = [
                    CanonicalDocumentElement(
                        id=uuid4(),
                        document_id=document_id,
                        ordinal=index,
                        element_type=type(element).__name__,
                        text=str(element).strip(),
                        locator=self._metadata_locator(element, name, index),
                    )
                    for index, element in enumerate(elements, start=1)
                    if str(element).strip()
                ]
                if normalized:
                    return normalized
            except Exception:  # noqa: BLE001, S110 - approved parser fallback boundary
                pass
            return self._native_fallback(document_id, name, suffix, content)
        finally:
            temporary_path.unlink(missing_ok=True)

    def _native_fallback(  # noqa: C901, PLR0912 - isolated format branches
        self,
        document_id: UUID,
        name: str,
        suffix: str,
        content: bytes,
    ) -> list[CanonicalDocumentElement]:
        if suffix == ".pdf":
            reader = PdfReader(BytesIO(content))
            if reader.is_encrypted:
                message = "encrypted_pdf_not_supported"
                raise ValueError(message)
            entries = [
                (page.extract_text() or "", {"file": name, "page": index})
                for index, page in enumerate(reader.pages, start=1)
            ]
        elif suffix == ".docx":
            with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as handle:
                handle.write(content)
                path = Path(handle.name)
            try:
                document = Document(str(path))
                entries = [
                    (paragraph.text, {"file": name, "paragraph": index})
                    for index, paragraph in enumerate(document.paragraphs, start=1)
                ]
                for table_index, table in enumerate(document.tables, start=1):
                    for row_index, row in enumerate(table.rows, start=1):
                        entries.append(
                            (
                                " | ".join(cell.text for cell in row.cells),
                                {"file": name, "table": table_index, "row": row_index},
                            )
                        )
            finally:
                path.unlink(missing_ok=True)
        elif suffix == ".pptx":
            presentation = Presentation(BytesIO(content))
            entries = []
            for slide_index, slide in enumerate(presentation.slides, start=1):
                for shape_index, shape in enumerate(slide.shapes, start=1):
                    text = getattr(shape, "text", "")
                    entries.append(
                        (
                            text,
                            {"file": name, "slide": slide_index, "shape": shape_index},
                        )
                    )
        elif suffix == ".xlsx":
            with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as handle:
                handle.write(content)
                path = Path(handle.name)
            try:
                workbook = load_workbook(path, read_only=True, data_only=True)
                entries = []
                for sheet in workbook.worksheets:
                    for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                        text = " | ".join(str(value) for value in row if value is not None)
                        entries.append(
                            (text, {"file": name, "sheet": sheet.title, "row": row_index})
                        )
            finally:
                path.unlink(missing_ok=True)
        else:
            decoded = content.decode("utf-8", errors="replace")
            if suffix == ".csv":
                rows = csv.reader(StringIO(decoded))
                entries = [
                    (" | ".join(row), {"file": name, "row": index})
                    for index, row in enumerate(rows, start=1)
                ]
            else:
                entries = [
                    (line, {"file": name, "line": index})
                    for index, line in enumerate(decoded.splitlines(), start=1)
                ]
        return [
            CanonicalDocumentElement(
                id=uuid4(),
                document_id=document_id,
                ordinal=index,
                element_type="NativeFallbackText",
                text=text.strip(),
                locator=locator,
            )
            for index, (text, locator) in enumerate(entries, start=1)
            if text.strip()
        ]

    @staticmethod
    def _metadata_locator(
        element: object,
        name: str,
        ordinal: int,
    ) -> dict[str, Any]:
        metadata = getattr(element, "metadata", None)
        values = metadata.to_dict() if metadata is not None else {}
        locator: dict[str, Any] = {"file": name, "element": ordinal}
        for source, target in (
            ("page_number", "page"),
            ("sheet_name", "sheet"),
            ("text_as_html", "has_table_html"),
        ):
            if values.get(source) is not None:
                locator[target] = True if source == "text_as_html" else values[source]
        return locator

    @staticmethod
    def _suffix_for(content_type: str) -> str:
        return {
            "application/pdf": ".pdf",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": ".xlsx",
            "application/vnd.openxmlformats-officedocument.presentationml.presentation": ".pptx",
            "text/csv": ".csv",
            "text/markdown": ".md",
        }.get(content_type, ".txt")
